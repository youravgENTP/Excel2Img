"""Border-geometry based table detection."""

from __future__ import annotations

from collections import deque
from dataclasses import dataclass
from typing import Iterable

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(frozen=True, order=True)
class TableRange:
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def coordinate(self) -> str:
        return (
            f"{get_column_letter(self.min_col)}{self.min_row}:"
            f"{get_column_letter(self.max_col)}{self.max_row}"
        )


def _visible(side: object) -> bool:
    style = getattr(side, "style", None)
    return bool(style and str(style).lower() != "none")


def _border_segments(ws: Worksheet) -> tuple[set[tuple[int, int]], set[tuple[int, int]]]:
    """Return horizontal (boundary-row, col) and vertical (row, boundary-col) units."""
    horizontal: set[tuple[int, int]] = set()
    vertical: set[tuple[int, int]] = set()

    # Only instantiated cells matter. This avoids scanning Excel's occasionally
    # enormous, stale used range. Loading a workbook materializes merged edges.
    for cell in ws._cells.values():  # openpyxl has no public sparse-cell iterator
        row, col = cell.row, cell.column
        border = cell.border
        if _visible(border.top):
            horizontal.add((row - 1, col))
        if _visible(border.bottom):
            horizontal.add((row, col))
        if _visible(border.left):
            vertical.add((row, col - 1))
        if _visible(border.right):
            vertical.add((row, col))
    return horizontal, vertical


def _neighbors(
    cell: tuple[int, int],
    horizontal: set[tuple[int, int]],
    vertical: set[tuple[int, int]],
) -> Iterable[tuple[int, int]]:
    row, col = cell
    if (row - 1, col) not in horizontal:
        yield row - 1, col
    if (row, col) not in horizontal:
        yield row + 1, col
    if (row, col - 1) not in vertical:
        yield row, col - 1
    if (row, col) not in vertical:
        yield row, col + 1


def detect_tables(ws: Worksheet) -> list[TableRange]:
    """Find maximal edge-adjacent cell regions enclosed by visible borders."""
    horizontal, vertical = _border_segments(ws)
    if not horizontal and not vertical:
        return []

    rows = [r for r, _ in horizontal] + [r for r, _ in vertical]
    cols = [c for _, c in horizontal] + [c for _, c in vertical]
    min_row, max_row = max(0, min(rows) - 1), max(rows) + 1
    min_col, max_col = max(0, min(cols) - 1), max(cols) + 1

    domain = {
        (row, col)
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    }
    outside: set[tuple[int, int]] = set()
    queue = deque([(min_row, min_col)])
    while queue:
        cell = queue.popleft()
        if cell in outside or cell not in domain:
            continue
        outside.add(cell)
        queue.extend(_neighbors(cell, horizontal, vertical))

    # Row/column zero are padding, not worksheet cells.
    enclosed = {cell for cell in domain - outside if cell[0] >= 1 and cell[1] >= 1}
    tables: list[TableRange] = []
    while enclosed:
        start = min(enclosed)
        component: set[tuple[int, int]] = set()
        queue = deque([start])
        while queue:
            cell = queue.popleft()
            if cell not in enclosed:
                continue
            enclosed.remove(cell)
            component.add(cell)
            row, col = cell
            queue.extend(((row - 1, col), (row + 1, col), (row, col - 1), (row, col + 1)))
        component_rows = [r for r, _ in component]
        component_cols = [c for _, c in component]
        tables.append(
            TableRange(
                min(component_rows), min(component_cols), max(component_rows), max(component_cols)
            )
        )
    return sorted(tables)


def workbook_tables(workbook: object) -> list[tuple[str, TableRange]]:
    found: list[tuple[str, TableRange]] = []
    for ws in workbook.worksheets:
        if not ws.title.startswith("!"):
            found.extend((ws.title, table) for table in detect_tables(ws))
    return found

