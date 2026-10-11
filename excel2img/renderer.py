"""High-fidelity rendering through Microsoft Excel for Mac."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import time
import uuid
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterator

import pymupdf

EXCEL_APP = Path("/Applications/Microsoft Excel.app")
EXCEL_TEMP_ROOT = (
    Path.home() / "Library" / "Containers" / "com.microsoft.Excel" / "Data" / "tmp" / "Excel2Img"
)
COMMAND_TIMEOUT = 30.0
RECOVERY_TIMEOUT = 10.0
SECURITY_VALUES = {
    "msoAutomationSecurityLow",
    "msoAutomationSecurityByUI",
    "msoAutomationSecurityForceDisable",
}


class RenderError(RuntimeError):
    """A workbook or range could not be rendered."""


class AutomationTimeout(RenderError):
    """Excel did not answer an automation request in time."""


class AutomationUnrecoverableError(RenderError):
    """Excel could not be returned to a known responsive state."""


@dataclass(frozen=True)
class ExcelSettings:
    automation_security: str
    display_alerts: bool


def _run_osascript(
    source: str,
    language: str = "AppleScript",
    *,
    timeout: float = COMMAND_TIMEOUT,
    phase: str = "Excel automation",
) -> str:
    command = ["/usr/bin/osascript"]
    if language != "AppleScript":
        command.extend(["-l", language])
    command.extend(["-e", source])
    try:
        completed = subprocess.run(
            command, text=True, capture_output=True, timeout=timeout, check=False
        )
    except subprocess.TimeoutExpired as exc:
        raise AutomationTimeout(f"{phase} timed out after {timeout:g} seconds") from exc
    if completed.returncode:
        detail = completed.stderr.strip() or completed.stdout.strip() or "unknown AppleScript error"
        raise RenderError(f"{phase} failed: {detail}")
    return completed.stdout.strip()


def _as_literal(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def check_excel_connection(timeout: float = RECOVERY_TIMEOUT) -> str:
    if not EXCEL_APP.exists():
        raise RenderError(f"Microsoft Excel was not found at {EXCEL_APP}.")
    return _run_osascript(
        'tell application "Microsoft Excel" to get name',
        timeout=timeout,
        phase="Excel connection",
    )


def read_excel_settings(timeout: float = RECOVERY_TIMEOUT) -> ExcelSettings:
    result = _run_osascript(
        '''
tell application "Microsoft Excel"
    set securityName to (automation security as text)
    set alertsName to (display alerts as text)
    return securityName & "|||" & alertsName
end tell
''',
        timeout=timeout,
        phase="Reading Excel application settings",
    )
    try:
        security, alerts = result.rsplit("|||", 1)
    except ValueError as exc:
        raise RenderError(f"Excel returned unrecognized settings: {result!r}") from exc
    if security not in SECURITY_VALUES or alerts.lower() not in {"true", "false"}:
        raise RenderError(f"Excel returned unrecognized settings: {result!r}")
    return ExcelSettings(security, alerts.lower() == "true")


def _set_automation_security(value: str, timeout: float = RECOVERY_TIMEOUT) -> None:
    if value not in SECURITY_VALUES:
        raise ValueError(f"Unsupported Excel automation security value: {value}")
    _run_osascript(
        f'tell application "Microsoft Excel" to set automation security to {value}',
        timeout=timeout,
        phase=f"Setting Excel automation security to {value}",
    )


def _open_workbook(path: Path, timeout: float = COMMAND_TIMEOUT) -> str:
    # Avoid an AppleScript error handler here: Excel's original error passes
    # through and no undefined workbook variable can mask an opening failure.
    script = f'''
set openedWorkbookName to ""
tell application "Microsoft Excel"
    set openedWorkbook to open workbook workbook file name {_as_literal(str(path))} update links do not update links read only true
    set openedWorkbookName to name of openedWorkbook
end tell
return openedWorkbookName
'''
    return _run_osascript(script, timeout=timeout, phase="Opening temporary workbook")


def _close_workbook_if_open(workbook_name: str, timeout: float = RECOVERY_TIMEOUT) -> bool:
    script = f'''
tell application "Microsoft Excel"
    if exists workbook {_as_literal(workbook_name)} then
        set temporaryWorkbook to workbook {_as_literal(workbook_name)}
        close temporaryWorkbook saving no
        return "closed"
    end if
end tell
return "not-open"
'''
    return (
        _run_osascript(script, timeout=timeout, phase="Closing temporary workbook") == "closed"
    )


def _raise_with_recovery(primary: BaseException, recovery_errors: list[str]) -> None:
    if not recovery_errors:
        raise primary
    detail = "; ".join(recovery_errors)
    raise AutomationUnrecoverableError(
        f"{primary}. Excel recovery also failed: {detail}. "
        "No user workbook was closed; resolve any Excel dialog and run --check."
    ) from primary


@contextmanager
def temporary_excel_workbook(
    source: Path,
    log: Callable[[str], None] = print,
    *,
    open_timeout: float = COMMAND_TIMEOUT,
) -> Iterator[str]:
    if not EXCEL_APP.exists():
        raise RenderError(f"Microsoft Excel was not found at {EXCEL_APP}.")

    log("Creating temporary workbook copy...")
    try:
        EXCEL_TEMP_ROOT.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise RenderError(
            f"Cannot create Excel's private temporary directory {EXCEL_TEMP_ROOT}: {exc}"
        ) from exc
    with tempfile.TemporaryDirectory(prefix="job-", dir=EXCEL_TEMP_ROOT) as directory:
        # Excel refuses duplicate filenames. The random name also identifies
        # the only workbook that cleanup is ever allowed to close.
        temporary_name = f"render-{uuid.uuid4().hex}{source.suffix.lower()}"
        temp_path = Path(directory) / temporary_name
        shutil.copy2(source, temp_path)

        original_security: str | None = None
        security_changed = False
        workbook_name: str | None = None
        primary: BaseException | None = None
        recovery_errors: list[str] = []

        # XLSX cannot contain VBA. Only XLSM needs a global security change, and
        # Python restores the exact observed value immediately after opening.
        if source.suffix.lower() == ".xlsm":
            original_security = read_excel_settings().automation_security
            if original_security != "msoAutomationSecurityForceDisable":
                security_changed = True

        try:
            if security_changed:
                _set_automation_security("msoAutomationSecurityForceDisable")
            log("Opening temporary workbook...")
            workbook_name = _open_workbook(temp_path, timeout=open_timeout)
        except BaseException as exc:
            primary = exc
        finally:
            if security_changed and original_security is not None:
                try:
                    _set_automation_security(original_security)
                except BaseException as exc:
                    recovery_errors.append(f"security restoration: {exc}")

        if primary is not None:
            # The killed osascript may leave Excel completing the open. Probe it,
            # then conditionally close only our random temporary workbook.
            if isinstance(primary, (AutomationTimeout, KeyboardInterrupt, SystemExit)):
                try:
                    check_excel_connection()
                    _close_workbook_if_open(temporary_name)
                except BaseException as exc:
                    recovery_errors.append(f"responsiveness check: {exc}")
            if isinstance(primary, (KeyboardInterrupt, SystemExit)):
                for error in recovery_errors:
                    log(f"Warning: Excel cleanup failed: {error}")
                raise primary
            _raise_with_recovery(primary, recovery_errors)

        if recovery_errors:
            if workbook_name:
                try:
                    _close_workbook_if_open(workbook_name)
                except BaseException as exc:
                    recovery_errors.append(f"temporary workbook cleanup: {exc}")
            _raise_with_recovery(
                RenderError("Temporary workbook opened, but Excel settings were not restored"),
                recovery_errors,
            )

        assert workbook_name is not None
        log("Temporary workbook opened.")
        render_error: BaseException | None = None
        try:
            yield workbook_name
        except BaseException as exc:
            render_error = exc
            raise
        finally:
            try:
                _close_workbook_if_open(workbook_name)
            except BaseException as exc:
                if render_error is None:
                    raise AutomationUnrecoverableError(
                        f"Rendering completed, but the temporary workbook could not be closed: {exc}. "
                        "No user workbook was closed."
                    ) from exc
                log(f"Warning: could not close temporary workbook: {exc}")


def _copy_range(workbook_name: str, sheet_name: str, coordinate: str) -> None:
    script = f'''
tell application "Microsoft Excel"
    set targetWorkbook to workbook {_as_literal(workbook_name)}
    set targetWorksheet to worksheet {_as_literal(sheet_name)} of targetWorkbook
    set targetRange to range {_as_literal(coordinate)} of targetWorksheet
    activate object targetWorkbook
    activate object targetWorksheet
    goto reference targetRange
    copy picture targetRange appearance printer format picture
end tell
'''
    _run_osascript(script, phase=f"Copying {sheet_name}!{coordinate} as a picture")
    time.sleep(0.15)


def _write_clipboard_pdf(destination: Path) -> None:
    script = f'''
ObjC.import('AppKit');
const pb = $.NSPasteboard.generalPasteboard;
let data = pb.dataForType('com.adobe.pdf');
if (!data) data = pb.dataForType('public.tiff');
if (!data) throw new Error('Excel did not place PDF or TIFF image data on the clipboard');
if (!data.writeToFileAtomically({_as_literal(str(destination))}, true))
    throw new Error('Could not write clipboard image data');
'''
    _run_osascript(
        script,
        language="JavaScript",
        phase="Reading the Excel picture from the macOS pasteboard",
    )


def _rasterize(source: Path, destination: Path, dpi: int) -> None:
    try:
        document = pymupdf.open(source)
        if document.page_count:
            page = document[0]
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(dpi / 72, dpi / 72), alpha=False
            )
            pixmap.save(destination)
            document.close()
            return
        document.close()
    except (pymupdf.FileDataError, RuntimeError, ValueError):
        pass

    try:
        from PIL import Image

        with Image.open(source) as image:
            image.save(destination, format="PNG", dpi=(dpi, dpi))
    except Exception as exc:
        raise RenderError(f"Could not convert Excel clipboard image: {exc}") from exc


class ExcelRenderer:
    def __init__(self, dpi: int = 300, log: Callable[[str], None] = print):
        self.dpi = dpi
        self.log = log

    def render_open_workbook(
        self, workbook_name: str, sheet_name: str, coordinate: str, destination: Path
    ) -> None:
        self.log(f"Rendering {sheet_name}!{coordinate}...")
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="excel2img-clipboard-") as directory:
            clipboard_data = Path(directory) / "range.data"
            _copy_range(workbook_name, sheet_name, coordinate)
            _write_clipboard_pdf(clipboard_data)
            _rasterize(clipboard_data, destination, self.dpi)


def automation_preflight(log: Callable[[str], None] = print) -> None:
    """Exercise the same open/copy/convert/close path used for real workbooks."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side

    log("Checking Microsoft Excel installation...")
    check_excel_connection()
    log("Excel connection: OK")
    settings = read_excel_settings()
    log(
        "Excel settings: automation security="
        f"{settings.automation_security}, display alerts={str(settings.display_alerts).lower()}"
    )
    if not settings.display_alerts:
        log("Warning: Excel display alerts is already false; Excel2Img will not change it.")

    with tempfile.TemporaryDirectory(prefix="excel2img-check-") as directory:
        source = Path(directory) / "preflight.xlsx"
        destination = Path(directory) / "preflight.png"
        workbook = Workbook()
        ws = workbook.active
        sheet_name = ws.title
        ws["A1"] = "Excel2Img"
        thin = Side(style="thin")
        ws["A1"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
        workbook.save(source)
        workbook.close()
        renderer = ExcelRenderer(log=log)
        with temporary_excel_workbook(source, log=log) as workbook_name:
            renderer.render_open_workbook(workbook_name, sheet_name, "A1:A1", destination)
        if not destination.is_file() or destination.stat().st_size == 0:
            raise RenderError("Preflight did not produce a PNG image.")
        log("Picture copy and PNG conversion: OK")
