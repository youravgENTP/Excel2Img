import json

from openpyxl import Workbook

from excel2img.paths import iter_workbooks, load_directories


def test_config_expands_deduplicates_and_warns_for_missing(tmp_path):
    parent = tmp_path / "자료"
    child = parent / "하위 폴더"
    child.mkdir(parents=True)
    missing_path = tmp_path / "없음"
    config = tmp_path / "config.json"
    config.write_text(
        json.dumps({"directories": [str(child), str(parent), str(parent), str(missing_path)]}),
        encoding="utf-8",
    )
    existing, missing = load_directories(config)
    assert existing == [parent.resolve()]
    assert missing == [missing_path.resolve()]


def test_recursive_iteration_excludes_generated_and_ignored(tmp_path):
    nested = tmp_path / "강의" / "1 주차"
    nested.mkdir(parents=True)
    Workbook().save(nested / "정상.xlsx")
    Workbook().save(nested / "!제외.xlsx")
    generated = nested / "정상_Img"
    generated.mkdir()
    Workbook().save(generated / "inside.xlsx")
    assert [path.name for path in iter_workbooks([tmp_path])] == ["정상.xlsx"]

