from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from excel2img.detector import TableRange, detect_tables, workbook_tables


THIN = Side(style="thin")


def outline(ws, min_row, min_col, max_row, max_col):
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.border = Border(
                top=THIN if row == min_row else cell.border.top,
                bottom=THIN if row == max_row else cell.border.bottom,
                left=THIN if col == min_col else cell.border.left,
                right=THIN if col == max_col else cell.border.right,
            )


def detected(*rectangles):
    workbook = Workbook()
    ws = workbook.active
    for rectangle in rectangles:
        outline(ws, *rectangle)
    return detect_tables(ws)


def test_single_rectangular_table():
    assert detected((2, 3, 5, 7)) == [TableRange(2, 3, 5, 7)]


def test_two_vertically_separated_tables():
    assert detected((1, 1, 2, 3), (5, 1, 7, 3)) == [
        TableRange(1, 1, 2, 3),
        TableRange(5, 1, 7, 3),
    ]


def test_two_horizontally_separated_tables():
    assert detected((2, 1, 4, 2), (2, 5, 4, 8)) == [
        TableRange(2, 1, 4, 2),
        TableRange(2, 5, 4, 8),
    ]


def test_diagonal_and_corner_touching_tables_stay_separate():
    assert detected((1, 1, 2, 2), (3, 3, 4, 4)) == [
        TableRange(1, 1, 2, 2),
        TableRange(3, 3, 4, 4),
    ]


def test_independent_tables_can_share_row_indices():
    assert len(detected((3, 1, 6, 2), (4, 5, 5, 7))) == 2


def test_independent_tables_can_share_column_indices():
    assert len(detected((1, 3, 2, 5), (5, 4, 7, 6))) == 2


def test_merged_header_survives_save_and_load(tmp_path):
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.merge_cells("B2:E2")
    outline(ws, 2, 2, 5, 5)
    workbook.save(path)
    loaded = load_workbook(path)
    assert detect_tables(loaded.active) == [TableRange(2, 2, 5, 5)]


def test_blank_interior_cells_are_included():
    assert detected((2, 2, 8, 9)) == [TableRange(2, 2, 8, 9)]


def test_outline_without_internal_gridlines_is_one_table():
    assert detected((1, 1, 10, 10)) == [TableRange(1, 1, 10, 10)]


def test_unbordered_notes_and_decorative_line_are_ignored():
    workbook = Workbook()
    ws = workbook.active
    ws["A1"] = "ordinary note"
    ws["Z20"] = "another note"
    outline(ws, 4, 3, 6, 5)
    ws["H2"].border = Border(bottom=THIN)
    assert detect_tables(ws) == [TableRange(4, 3, 6, 5)]


def test_ignored_worksheet():
    workbook = Workbook()
    visible = workbook.active
    visible.title = "약물"
    outline(visible, 1, 1, 2, 2)
    ignored = workbook.create_sheet("!Scratch")
    outline(ignored, 1, 1, 3, 3)
    assert workbook_tables(workbook) == [("약물", TableRange(1, 1, 2, 2))]

